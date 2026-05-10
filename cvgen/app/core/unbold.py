"""Post-render formatting fix — labels (text before colon) bold, values regular."""
from pathlib import Path

from docx import Document


def apply_label_bold_rule(docx_path: Path):
    """
    For each paragraph in the rendered document:
    - Text before and including the first ':' → bold (label)
    - Text after the ':' → regular weight (value)
    - Paragraphs without ':' → unchanged

    Handles the case where docxtpl loops produce paragraphs whose formatting
    has been reset to the paragraph default.
    """
    doc = Document(str(docx_path))

    def process_paragraph(p):
        full_text = "".join(r.text for r in p.runs)
        if ":" not in full_text:
            return
        colon_idx = full_text.index(":") + 1

        pos = 0
        for run in p.runs:
            run_len = len(run.text)
            run_start = pos
            run_end = pos + run_len
            pos = run_end

            if run_end <= colon_idx:
                run.bold = True
            elif run_start >= colon_idx:
                run.bold = False
            else:
                # Run straddles colon — split
                split_at = colon_idx - run_start
                label_part = run.text[:split_at]
                value_part = run.text[split_at:]
                run.text = label_part
                run.bold = True
                new_run = p.add_run(value_part)
                new_run.bold = False
                if run.font.name:
                    new_run.font.name = run.font.name
                if run.font.size:
                    new_run.font.size = run.font.size
                run._r.addnext(new_run._r)

    for p in doc.paragraphs:
        process_paragraph(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph(p)

    doc.save(str(docx_path))
